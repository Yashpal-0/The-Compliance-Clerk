import unittest
import tempfile
import os
from pathlib import Path
import openpyxl

from src.reporter.reporter import write_excel, write_failed_log
from src.linker.linker import link_records

class TestReporter(unittest.TestCase):

    def setUp(self):
        # Create a temporary directory for output files
        self.temp_dir = tempfile.TemporaryDirectory()
        self.output_dir = self.temp_dir.name

    def tearDown(self):
        # Clean up the temporary directory
        self.temp_dir.cleanup()

    def test_write_excel_empty_rows(self):
        output_path = os.path.join(self.output_dir, "empty.xlsx")
        
        # Test with empty rows
        write_excel([], output_path)
        
        self.assertTrue(os.path.exists(output_path))
        
        # Verify sheets and headers exist
        wb = openpyxl.load_workbook(output_path)
        self.assertIn("Sheet1", wb.sheetnames)
        
        # Verify NA PDF empty sheet
        ws_na = wb["Sheet1"]
        self.assertEqual(ws_na.max_row, 1)
        self.assertEqual(ws_na.cell(row=1, column=1).value, "Sr.no.")

    def test_write_excel_with_data(self):
        output_path = os.path.join(self.output_dir, "data.xlsx")
        
        rows = [
            {
                "doc_type": "na_order",
                "village": "TestVillage",
                "survey_number": "123",
                "land_area_sqm": "1000",
                "source_file": "doc1.pdf"
            },
            {
                "doc_type": "lease_deed",
                "dnr_number": "DNR999",
                "stamp_duty": "10000",
                "source_file": "doc2.pdf"
            }
        ]
        
        linked_rows = link_records(rows)
        write_excel(linked_rows, output_path)
        
        self.assertTrue(os.path.exists(output_path))
        
        wb = openpyxl.load_workbook(output_path)
        
        ws = wb["Sheet1"]
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(row=2, column=2).value, "TestVillage")
        self.assertEqual(ws.cell(row=2, column=3).value, "123")
        self.assertEqual(ws.cell(row=2, column=4).value, "1000")
        
        self.assertEqual(ws.cell(row=3, column=7).value, "DNR999")
        
    def test_write_failed_log_empty(self):
        # Empty array means no file should be created
        write_failed_log([], self.output_dir)
        log_path = Path(self.output_dir) / "failed_extractions.txt"
        self.assertFalse(log_path.exists())

    def test_write_failed_log_with_data(self):
        failed_list = [
            {"file": "missing.pdf", "reason": "No text found"},
            {"file": "corrupt.pdf", "reason": "Decryption error"}
        ]
        
        write_failed_log(failed_list, self.output_dir)
        log_path = Path(self.output_dir) / "failed_extractions.txt"
        self.assertTrue(log_path.exists())
        
        with open(log_path, "r") as f:
            content = f.read()
            
        self.assertIn("missing.pdf: No text found", content)
        self.assertIn("corrupt.pdf: Decryption error", content)

if __name__ == '__main__':
    unittest.main()
